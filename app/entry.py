"""
Data-entry and live-stream routes.

Split out from main.py so the read API stays easy to scan: everything that
changes the database, plus the event stream the dashboard listens on, lives here
and is mounted onto the app as a router.

Supports dual-mode execution: writes directly to PostgreSQL when configured,
and maintains in-memory and client-side reactive state when running in serverless/offline mode.
"""

from __future__ import annotations

import io
import json
import logging
import os
import threading
import time
import uuid
from datetime import date, datetime
from fastapi import APIRouter, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import StreamingResponse
from starlette.concurrency import run_in_threadpool
import openpyxl
import psycopg
from psycopg.errors import IntegrityError

from . import fallback
from .db import (is_db_ready, fetch_all, pool, DSN, ensure_pool_open,
                 connect as db_connect, session)
from .events import broker
from etl.dimensions import activate_period, report_period
from etl.load_dsr import Loader, workbook_problem
from .write import (AllotmentIn, BookingIn, BookingPatch, LeadIn, RegistrationIn,
                    TestDriveIn, VehicleIn, create_allotment, create_booking,
                    create_lead, create_registration, create_test_drive,
                    delete_row, update_booking, upsert_vehicle)

log = logging.getLogger("dsr.entry")
router = APIRouter()

# Everything a reporting month owns. Facts are keyed by the load that produced
# them (load_period_id); lead and booking additionally carry a business
# period_id, which is where hand-entered rows are attached.
PERIOD_FACTS = ["registration", "allotment", "booking", "test_drive", "lead"]
PERIOD_TARGET_TABLES = [
    "target_daily_tracker", "target_booking_commitment",
    "target_channel_funnel", "target_consultant_scorecard",
]

# Month number -> the three-letter form the period labels use (AUG2026).
_MONTH_ABBR = {1: "JAN", 2: "FEB", 3: "MAR", 4: "APR", 5: "MAY", 6: "JUN",
               7: "JUL", 8: "AUG", 9: "SEP", 10: "OCT", 11: "NOV", 12: "DEC"}

def _resolve_period(period: str | None, filename: str) -> tuple[str, date, date]:
    """
    Work out which month a report is for - or refuse to guess.

    This used to default to August 2026 whenever it could not tell. A file
    called "Test DSR.xlsx" therefore loaded silently into AUG2026 and replaced
    that month's bookings, test drives, allotments and registrations. Getting
    this wrong destroys a month of data, so an unanswerable case is now an
    error the uploader can act on rather than a guess nobody sees.

    FIX (2026-10-08): the rule itself moved to etl.dimensions.report_period so
    the command-line loader - and through it the Drive and email sync daemons,
    which always loaded into AUG2026 - applies exactly the same one. This is
    now only the HTTP wrapper around it.
    """
    try:
        return report_period(period, filename)
    except ValueError as exc:
        raise HTTPException(400, str(exc)) from exc


# =====================================================================
# Live change stream
# =====================================================================

@router.get("/api/events", tags=["live"], include_in_schema=False)
async def events():
    """
    Server-sent events: one `change` frame whenever the database moves.
    """
    if os.environ.get("VERCEL") or not is_db_ready():
        async def vercel_stream():
            yield 'event: ready\ndata: {"connected": false, "serverless": true}\n\n'
        return StreamingResponse(
            vercel_stream(),
            media_type="text/event-stream",
            headers={"Cache-Control": "no-cache"},
        )

    return StreamingResponse(
        broker.stream(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache, no-transform",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@router.get("/api/live-status", tags=["live"])
def live_status():
    """Whether the listener is attached, and how many dashboards are watching."""
    if os.environ.get("VERCEL") or not is_db_ready():
        return {"connected": False, "serverless": True, "subscribers": 0, "events_seen": 0, "last_error": None}
    return broker.status


# =====================================================================
# Writes
# =====================================================================

def _commit(fn, *args):
    """Run one writer inside a transaction and turn its errors into HTTP codes."""
    ensure_pool_open()
    try:
        with pool.connection(timeout=10.0) as cx:
            with cx.transaction():
                return fn(cx, *args)
    except HTTPException:
        raise
    except PermissionError as exc:
        raise HTTPException(409, str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(422, str(exc)) from exc
    except IntegrityError as exc:
        raise HTTPException(409, str(exc).strip().splitlines()[0]) from exc
    except Exception as exc:
        log.exception("Database transaction failed: %s", exc)
        raise HTTPException(500, f"Database transaction failed: {exc}") from exc


@router.post("/api/leads", status_code=201, tags=["entry"])
def add_lead(body: LeadIn):
    if is_db_ready():
        res = _commit(create_lead, body)
        broker.notify_sync("lead")
        return res
    return fallback.store.add_lead(body.model_dump())


@router.post("/api/bookings", status_code=201, tags=["entry"])
def add_booking(body: BookingIn):
    if is_db_ready():
        res = _commit(create_booking, body)
        broker.notify_sync("booking")
        return res
    return fallback.store.add_booking(body.model_dump())


@router.patch("/api/bookings/{booking_id}", tags=["entry"])
def patch_booking(booking_id: int, body: BookingPatch):
    if is_db_ready():
        res = _commit(update_booking, booking_id, body)
        broker.notify_sync("booking")
        return res
    return fallback.store.patch_booking(booking_id, body.model_dump(exclude_unset=True))


@router.post("/api/test-drives", status_code=201, tags=["entry"])
def add_test_drive(body: TestDriveIn):
    if is_db_ready():
        res = _commit(create_test_drive, body)
        broker.notify_sync("test_drive")
        return res
    return fallback.store.add_test_drive(body.model_dump())


@router.post("/api/vehicles", status_code=201, tags=["entry"])
def add_vehicle(body: VehicleIn):
    if is_db_ready():
        res = _commit(upsert_vehicle, body)
        broker.notify_sync("vehicle")
        return res
    return fallback.store.add_vehicle(body.model_dump())


@router.post("/api/allotments", status_code=201, tags=["entry"])
def add_allotment(body: AllotmentIn):
    if is_db_ready():
        res = _commit(create_allotment, body)
        broker.notify_sync("allotment")
        return res
    return fallback.store.add_allotment(body.model_dump())


@router.post("/api/registrations", status_code=201, tags=["entry"])
def add_registration(body: RegistrationIn):
    if is_db_ready():
        res = _commit(create_registration, body)
        broker.notify_sync("registration")
        return res
    return fallback.store.add_registration(body.model_dump())


# FIX (2026-10-08): every call to this route returned 500 - it passed
# (table, row_id) to delete_row(), which wanted (table, pk_column, row_id). The
# writer now resolves the key column itself (app/write.py). Two smaller things
# came with it: the entity is accepted the way the README and the offline store
# write it ("bookings", "test-drive"), and a row that does not exist is a 404
# rather than a 200 with `false` in the body.
@router.delete("/api/entries/{table}/{row_id}", status_code=200, tags=["entry"])
def delete_entry(table: str, row_id: int):
    entity = table.lower().replace("-", "_").rstrip("s")
    if is_db_ready():
        deleted = _commit(delete_row, entity, row_id)
        if not deleted:
            raise HTTPException(404, f"No {entity} with id {row_id}.")
        broker.notify_sync(entity)
        return {"deleted": True, "table": entity, "id": row_id}
    return fallback.store.delete_entry(table, row_id)


# =====================================================================
# Dropdowns and datalists for the entry drawer
# =====================================================================

@router.get("/api/entry-options", tags=["entry"])
def entry_options():
    if is_db_ready():
        try:
            # Seven queries over one checked-out connection rather than seven:
            # against a remote database the checkout costs more than the query.
            with session() as cx:
                q = lambda sql: cx.execute(sql).fetchall()
                return {
                    "consultants": [r["display_name"] for r in q(
                        "SELECT display_name FROM dim_consultant WHERE is_active "
                        "ORDER BY display_name")],
                    "sources": [r["name"] for r in q(
                        "SELECT name FROM dim_lead_source ORDER BY name")],
                    "models": [r["name"] for r in q(
                        "SELECT name FROM dim_model ORDER BY name")],
                    "variants": q(
                        "SELECT m.name AS model, v.name AS variant, v.long_model_text "
                        "FROM dim_variant v JOIN dim_model m USING (model_id) "
                        "ORDER BY m.name, v.name"),
                    "colours": [r["name"] for r in q(
                        "SELECT name FROM dim_colour ORDER BY name")],
                    "fulfilment_statuses": ["BOOKED", "NO_STOCK", "ALLOTED",
                                            "RETAILED", "CANCELLED"],
                    "open_bookings": q("""
                        SELECT b.booking_id, b.customer_name, m.name AS model,
                               dv.name AS variant, c.display_name AS consultant
                        FROM booking b
                        LEFT JOIN dim_model m      ON m.model_id = b.model_id
                        LEFT JOIN dim_variant dv   ON dv.variant_id = b.variant_id
                        LEFT JOIN dim_consultant c ON c.consultant_id = b.consultant_id
                        WHERE b.fulfilment_status IN ('BOOKED', 'NO_STOCK')
                        ORDER BY b.booking_date DESC NULLS LAST
                        LIMIT 200
                    """),
                    "free_chassis": q("""
                        SELECT chassis_number, model, variant, colour, stock_aging_days
                        FROM v_stock WHERE stock_status = 'FREESTOCK'
                        ORDER BY stock_aging_days DESC NULLS LAST
                    """),
                }
        except Exception:
            pass
    return fallback.get_entry_options()


@router.get("/api/recent-activity", tags=["entry"])
def recent_activity(limit: int = Query(25, le=100)):
    if is_db_ready():
        try:
            return fetch_all("""
                SELECT 'booking' AS kind, booking_id AS id, customer_name AS who,
                       loaded_at, entered_by
                  FROM booking WHERE origin = 'MANUAL'
                UNION ALL
                SELECT 'lead', lead_id, lead_name, loaded_at, entered_by
                  FROM lead WHERE origin = 'MANUAL'
                UNION ALL
                SELECT 'test drive', test_drive_id, lead_name, loaded_at, entered_by
                  FROM test_drive WHERE origin = 'MANUAL'
                UNION ALL
                SELECT 'allotment', allotment_id, customer_name, loaded_at, entered_by
                  FROM allotment WHERE origin = 'MANUAL'
                UNION ALL
                SELECT 'registration', registration_id, customer_name, loaded_at, entered_by
                  FROM registration WHERE origin = 'MANUAL'
                UNION ALL
                SELECT 'vehicle', vehicle_id, chassis_number, loaded_at, entered_by
                  FROM vehicle WHERE origin = 'MANUAL'
                ORDER BY loaded_at DESC
                LIMIT %s
            """, (limit,))
        except Exception:
            pass
    return fallback.store.get_recent_activity(limit)


# =====================================================================
# Reporting period
# =====================================================================

@router.get("/api/periods", tags=["entry"])
def periods():
    if is_db_ready():
        try:
            res = fetch_all("""
                SELECT p.label, p.period_start, p.period_end, p.is_active,
                       (SELECT count(*) FROM lead l    WHERE l.period_id = p.period_id) AS leads,
                       (SELECT count(*) FROM booking b WHERE b.period_id = p.period_id) AS bookings
                FROM dim_period p
                ORDER BY p.period_start DESC
            """)
            if res:
                return res
        except Exception as e:
            import traceback
            traceback.print_exc()
            raise
    return fallback.get_periods()


@router.post("/api/period/{label}/activate", tags=["entry"])
def set_active_period(label: str):
    if is_db_ready():
        try:
            return _commit(activate_period, label)
        except Exception as e:
            import traceback
            traceback.print_exc()
            raise HTTPException(status_code=500, detail=str(e))
    return {"activated": label}


@router.get("/api/periods/{label}/contents", tags=["entry"])
def period_contents(label: str):
    """
    What deleting this month would remove, so the confirmation can say so
    rather than asking the user to take it on faith.
    """
    if not is_db_ready():
        raise HTTPException(503, "The database is not reachable.")
    with session() as cx:
        row = cx.execute(
            "SELECT period_id, label, is_active FROM dim_period WHERE upper(label) = upper(%s)",
            (label,)).fetchone()
        if not row:
            raise HTTPException(404, f"There is no month called {label}.")
        pid = row["period_id"]
        counts = {}
        for table in PERIOD_FACTS:
            counts[table] = cx.execute(
                f"SELECT count(*) AS n FROM {table} "
                f"WHERE load_period_id = %s OR period_id = %s"
                if table in ("lead", "booking") else
                f"SELECT count(*) AS n FROM {table} WHERE load_period_id = %s",
                (pid, pid) if table in ("lead", "booking") else (pid,)).fetchone()["n"]
        manual = cx.execute(
            "SELECT count(*) AS n FROM lead WHERE origin = 'MANUAL' AND period_id = %s",
            (pid,)).fetchone()["n"] + cx.execute(
            "SELECT count(*) AS n FROM booking WHERE origin = 'MANUAL' AND period_id = %s",
            (pid,)).fetchone()["n"]
        # What uploading a workbook in its place removes - less than deleting
        # the month: a replace clears only what a workbook loaded into it, and
        # keeps the rows typed into the dashboard (see Loader.reset). The
        # upload panel showed the delete figures and said the hand-entered rows
        # would be replaced too; they never were.
        replaces = {}
        for table in PERIOD_FACTS:
            n = cx.execute(f"SELECT count(*) AS n FROM {table} "
                           f"WHERE origin = 'WORKBOOK' AND load_period_id = %s", (pid,)).fetchone()["n"]
            if n:
                replaces[table] = n
        return {
            "label": row["label"],
            "is_active": row["is_active"],
            "counts": {k: v for k, v in counts.items() if v},
            "total": sum(counts.values()),
            "hand_entered": manual,
            "replaces": replaces,
            "replaces_total": sum(replaces.values()),
        }


@router.post("/api/periods/{label}/clear", tags=["entry"])
def clear_period(label: str, confirm: str = Query(..., description="Must repeat the month's label")):
    """
    Empty a month without removing it.

    Distinct from deleting the month outright: the month, its date range and its
    place in the switcher survive, so the dashboard keeps working and shows
    zeros until a workbook is uploaded for it. This is the "start this month
    again" button, where delete is "this month should not exist".

    Uploading a workbook already replaces the month it belongs to, so this is
    only needed when someone wants the month emptied WITHOUT loading anything
    in its place.
    """
    if not is_db_ready():
        raise HTTPException(503, "The database is not reachable.")
    if confirm.strip().upper() != label.strip().upper():
        raise HTTPException(400, "Type the month's label exactly to confirm.")

    removed: dict[str, int] = {}
    with db_connect() as cx:
        with cx.transaction():
            row = cx.execute(
                "SELECT period_id, label FROM dim_period WHERE upper(label) = upper(%s) FOR UPDATE",
                (label,)).fetchone()
            if not row:
                raise HTTPException(404, f"There is no month called {label}.")
            pid, real_label = row[0], row[1]

            for table in PERIOD_FACTS:
                if table in ("lead", "booking"):
                    n = cx.execute(
                        f"DELETE FROM {table} WHERE load_period_id = %s OR period_id = %s",
                        (pid, pid)).rowcount
                else:
                    n = cx.execute(
                        f"DELETE FROM {table} WHERE load_period_id = %s", (pid,)).rowcount
                if n:
                    removed[table] = n
            for table in PERIOD_TARGET_TABLES:
                n = cx.execute(f"DELETE FROM {table} WHERE period_id = %s", (pid,)).rowcount
                if n:
                    removed[table] = n

    broker.notify_sync("lead")
    broker.notify_sync("booking")
    return {"cleared": real_label, "removed": removed, "total": sum(removed.values())}


@router.delete("/api/periods/{label}", tags=["entry"])
def delete_period(label: str, confirm: str = Query(..., description="Must repeat the month's label")):
    """
    Remove a reporting month and everything loaded under it.

    This cannot be undone from the dashboard - the only way back is to upload
    that month's workbook again - so it is guarded three ways: the caller has to
    repeat the label, the month must not be the active one, and the last
    remaining month cannot be removed.
    """
    if not is_db_ready():
        raise HTTPException(503, "The database is not reachable.")

    if confirm.strip().upper() != label.strip().upper():
        raise HTTPException(400, "Type the month's label exactly to confirm the deletion.")

    removed: dict[str, int] = {}
    with db_connect(autocommit=False) as cx:
        with cx.transaction():
            row = cx.execute(
                "SELECT period_id, label, is_active FROM dim_period "
                "WHERE upper(label) = upper(%s) FOR UPDATE",
                (label,)).fetchone()
            if not row:
                raise HTTPException(404, f"There is no month called {label}.")
            pid, real_label, is_active = row[0], row[1], row[2]

            if is_active:
                raise HTTPException(
                    409,
                    f"{real_label} is the month the dashboard is currently reporting on. "
                    f"Switch to another month first, then delete it.")

            if cx.execute("SELECT count(*) FROM dim_period").fetchone()[0] <= 1:
                raise HTTPException(409, "This is the only month on record; it cannot be removed.")

            # Facts first, then the targets, then the month itself, so no foreign
            # key is left pointing at a row that has gone.
            for table in PERIOD_FACTS:
                if table in ("lead", "booking"):
                    n = cx.execute(
                        f"DELETE FROM {table} WHERE load_period_id = %s OR period_id = %s",
                        (pid, pid)).rowcount
                else:
                    n = cx.execute(
                        f"DELETE FROM {table} WHERE load_period_id = %s", (pid,)).rowcount
                if n:
                    removed[table] = n
            for table in PERIOD_TARGET_TABLES:
                n = cx.execute(f"DELETE FROM {table} WHERE period_id = %s", (pid,)).rowcount
                if n:
                    removed[table] = n
            cx.execute("DELETE FROM etl_run WHERE notes LIKE %s", (f"%[{real_label}]%",))
            cx.execute("DELETE FROM dim_period WHERE period_id = %s", (pid,))

    broker.notify_sync("lead")
    broker.notify_sync("booking")
    return {
        "deleted": real_label,
        "removed": removed,
        "total": sum(removed.values()),
    }


# =====================================================================
# Excel Workbook Ingestion Pipeline
# =====================================================================

# =====================================================================
# Background ingestion
# =====================================================================
#
# Ingesting a DSR workbook takes ~100 seconds: openpyxl parses 22 sheets, then
# ~2,600 rows are written to a database a round trip away. Render's gateway
# ends any request at about 60 seconds, and a dev-server proxy defaults to the
# same, so the browser was told 502 while the ingest carried on and finished.
# The month loaded, the dashboard eventually showed it, and the person who
# pressed the button was told it had failed - so they pressed it again.
#
# No timeout can be configured away on a managed host, so the request no longer
# waits: it hands the work to a thread and returns a job id at once. Nothing in
# the path can time out a request that answers in a few milliseconds.
#
# The registry is a plain dict. Jobs are ephemeral and this runs as one process
# per service; a job lost to a restart is a job whose upload has to be redone,
# which is true of any in-flight request anyway.

_JOBS: dict[str, dict] = {}
_JOBS_LOCK = threading.Lock()
_JOB_TTL_SECONDS = 3600


def _job_set(job_id: str, **fields) -> None:
    with _JOBS_LOCK:
        job = _JOBS.setdefault(job_id, {})
        job.update(fields)
        job["updated_at"] = time.time()


def _job_get(job_id: str) -> dict | None:
    with _JOBS_LOCK:
        # Opportunistic sweep, so a long-lived process does not accumulate
        # every workbook anyone has ever uploaded.
        cutoff = time.time() - _JOB_TTL_SECONDS
        for k in [k for k, v in _JOBS.items()
                  if v.get("state") in ("done", "failed") and v.get("updated_at", 0) < cutoff]:
            _JOBS.pop(k, None)
        job = _JOBS.get(job_id)
        return dict(job) if job else None


def _upload_problem(content: bytes) -> str | None:
    """Why an uploaded file cannot be loaded as a DSR workbook, or None. Reads
    the tab names only, so it answers in a moment."""
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True)
    except Exception as exc:
        return f"Could not open the file as an Excel workbook ({exc}). Save it as .xlsx and try again."
    try:
        return workbook_problem(wb)
    finally:
        wb.close()


# FIX (2026-10-09): an upload died at its last step with
#   FATAL: (ECIRCUITBREAKER) too many authentication failures, new connections
#   are temporarily blocked
# after two minutes of parsing, and the person had to start again. Supabase's
# pooler blocks new connections from an address for a few minutes after it has
# seen repeated failed logins - and on Render that address is shared with other
# services (a second service with a stale DATABASE_URL keeps retrying and keeps
# the block open for everyone). The block clears by itself, so the load waits
# for it instead of failing. ONLY the circuit breaker is retried: any other
# connection error (a wrong password, a database that is down) fails at once,
# because retrying a failed login would feed the very block being waited out.
_BREAKER_TRIES = 6
_BREAKER_WAIT_SECONDS = 30


def _connect_for_load(job_id: str):
    """A direct connection for a workbook load, waiting out a pooler block."""
    for attempt in range(1, _BREAKER_TRIES + 1):
        try:
            return db_connect(autocommit=True)
        except psycopg.OperationalError as exc:
            if "ECIRCUITBREAKER" not in str(exc) or attempt == _BREAKER_TRIES:
                raise
            log.warning("ingest job %s: database temporarily blocked, retry %d/%d",
                        job_id, attempt, _BREAKER_TRIES - 1)
            _job_set(job_id, step=("waiting for the database - it is temporarily "
                                   f"blocking new connections (retry {attempt} of "
                                   f"{_BREAKER_TRIES - 1})"))
            time.sleep(_BREAKER_WAIT_SECONDS)


def _ingest_worker(job_id: str, content: bytes, fname: str, period_label: str,
                   period_start: date, period_end: date, uploaded_by: str,
                   mode: str = "replace") -> None:
    """Runs off the request thread. Only ever writes its result into _JOBS."""
    try:
        _job_set(job_id, state="running", step="parsing the workbook")
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)

        _job_set(job_id, step="writing to the database")
        with _connect_for_load(job_id) as cx:
            # One transaction around the whole load, so a workbook that fails
            # part way leaves the month exactly as it was.
            #
            # This ran on the autocommit connection alone until a replace of
            # AUG2026 died on a missing worksheet: reset() had already deleted
            # and committed the month's 2,154 leads, nothing was loaded in their
            # place, and the run failed before it could even write its etl_run
            # row - so the month was emptied and there was no record of what
            # had done it. A replace is a delete followed by a load, and the two
            # halves must not be separable.
            with cx.transaction():
                loader = Loader(cx, wb, period_label, period_start, period_end, mode=mode)
                counts = loader.run()
                cx.execute("""
                    INSERT INTO etl_run (source_file, file_modified, finished_at, row_counts, notes)
                    VALUES (%s, now(), now(), %s, %s)
                """, (
                    fname, json.dumps(counts),
                    f"Uploaded by {uploaded_by} ({mode})"
                    + (("; " + "; ".join(loader.warnings)) if loader.warnings else ""),
                ))

        # Tell every open dashboard, the same way a synchronous upload did.
        try:
            broker.notify_sync("etl_run")
        except Exception:
            pass

        _job_set(job_id, state="done", step="complete", counts=counts,
                 warnings=loader.warnings,
                 result={
                     "status": "success",
                     "message": (f"Added '{fname}' to {period_label}."
                                 if mode == "append" else
                                 f"Successfully ingested '{fname}' for {period_label}."),
                     "mode": mode,
                     "filename": fname,
                     "file_type": "excel",
                     "period": period_label,
                     "period_range": f"{period_start} to {period_end}",
                     "uploaded_by": uploaded_by,
                     "counts": counts,
                     "warnings": loader.warnings,
                     "environment": "postgres",
                 })
    except Exception as exc:
        log.exception("ingest job %s failed", job_id)
        _job_set(job_id, state="failed", step="failed", error=str(exc))


@router.post("/api/upload-report/start", tags=["ingestion"], status_code=202)
async def start_upload(
    file: UploadFile = File(..., description="DSR workbook (.xlsx, .xlsm)"),
    period: str | None = Form(None),
    uploaded_by: str = Form("Reporting Agent"),
    mode: str = Form("replace", description="replace the month, or append to it"),
    covers: str = Form("month", description="what the file covers: day, week or month"),
    covers_date: str | None = Form(None, description="the day, or the first day of the week (YYYY-MM-DD)"),
    covers_date_end: str | None = Form(None, description="the last day of the week (YYYY-MM-DD)"),
):
    """
    Accept a workbook and ingest it in the background.

    Returns at once with a job id; poll /api/upload-report/status/{job_id}.
    This is what the dashboard uses, because a ~100 second ingest cannot
    survive a 60 second gateway.
    """
    fname = file.filename or "unknown_report.xlsx"
    if not fname.lower().endswith((".xlsx", ".xlsm", ".xls")):
        raise HTTPException(400, "Background ingest takes an Excel workbook "
                                 "(.xlsx, .xlsm). Use /api/upload-report for CSV or text.")
    content = await file.read()
    if not content:
        raise HTTPException(400, "Uploaded file is empty.")
    # Before anything is queued - and so before a replace touches the month.
    problem = _upload_problem(content)
    if problem:
        raise HTTPException(400, problem)
    if not is_db_ready():
        raise HTTPException(503, "The database is not reachable.")

    mode = mode if mode in ("replace", "append") else "replace"
    covers = covers if covers in ("day", "week", "month") else "month"

    # A day or a week is still filed under the month it falls in: dim_period is
    # monthly, and every view, target and KPI is scoped that way. What the
    # choice changes is which month is picked (from the date rather than from a
    # label) and that appending is the only sane pairing - replacing a whole
    # month with one day's file would delete the rest of the month.
    if covers in ("day", "week"):
        if not covers_date:
            raise HTTPException(400, f"Uploading a {covers} needs the date it covers.")
        try:
            on = date.fromisoformat(covers_date)
        except ValueError:
            raise HTTPException(400, f"'{covers_date}' is not a date (expected YYYY-MM-DD).")

        covers_to = on
        if covers == "week" and covers_date_end:
            try:
                covers_to = date.fromisoformat(covers_date_end)
            except ValueError:
                raise HTTPException(400, f"'{covers_date_end}' is not a date (expected YYYY-MM-DD).")
            if covers_to < on:
                raise HTTPException(400, "The week's end date is before its start date.")

        # A week that straddles a month boundary is filed under the month it
        # STARTS in. Splitting it across two periods would mean two partial
        # loads whose targets and scorecards each belong to neither month.
        period_label, period_start, period_end = _resolve_period(
            f"{_MONTH_ABBR[on.month]}{on.year}", fname)
        mode = "append"
    else:
        period_label, period_start, period_end = _resolve_period(period, fname)

    job_id = uuid.uuid4().hex
    _job_set(job_id, state="queued", step="queued", filename=fname,
             period=period_label, started_at=time.time(), mode=mode, covers=covers,
             covers_from=covers_date, covers_to=(covers_date_end or covers_date))
    worker_args = (job_id, content, fname, period_label, period_start, period_end,
                   uploaded_by, mode)
    reply = {"job_id": job_id, "state": "queued", "period": period_label,
             "filename": fname, "mode": mode, "covers": covers,
             "covers_from": covers_date, "covers_to": covers_date_end or covers_date}

    # FIX (2026-10-09): on Vercel an upload never finished - the dashboard sat
    # on "Ingesting the workbook... 130s, still working" for ever. Vercel runs
    # the app as serverless functions and FREEZES a function the moment it has
    # sent its reply, so the background thread below stopped part-way, and the
    # job registry (_JOBS, in this process's memory) is not shared between
    # function instances, so the status polls could not find it either. On
    # Vercel the work is therefore done INSIDE this request, and the finished
    # result goes back in the reply itself (the page uses it without polling:
    # uploadWorkbookInBackground in frontend/src/api/client.js). vercel.json
    # pins the function to the Sydney region, next to the Supabase database, so
    # the many small queries of a load are fast enough to finish in the
    # function's time limit. Persistent hosts (Render, a laptop) keep the
    # background thread: the ~100 s load would not survive a 60 s gateway.
    if os.environ.get("VERCEL"):
        await run_in_threadpool(_ingest_worker, *worker_args)
        job = _job_get(job_id) or {}
        reply.update(state=job.get("state", "failed"), result=job.get("result"),
                     error=job.get("error"), counts=job.get("counts"),
                     warnings=job.get("warnings"))
        return reply

    threading.Thread(
        target=_ingest_worker,
        args=worker_args,
        daemon=True,
        name=f"ingest-{job_id[:8]}",
    ).start()

    return reply


@router.get("/api/upload-report/status/{job_id}", tags=["ingestion"])
def upload_status(job_id: str):
    """Where a background ingest has got to. 404 once the job has expired."""
    job = _job_get(job_id)
    if not job:
        raise HTTPException(404, "No such ingest job (it may have expired or the server restarted).")
    out = {k: v for k, v in job.items() if k != "updated_at"}
    if job.get("started_at"):
        out["elapsed"] = round(time.time() - job["started_at"], 1)
    return out


# FIX (2026-10-08): an Excel file sent here always failed with 500 "name 'mode'
# is not defined" - the Excel branch passed `mode` to the Loader but the route
# never declared it. It is now a form field (default "replace", the old
# intended behaviour), passed through for text files too, and echoed in the
# reply so the dashboard can say "added to" rather than "replaced".
@router.post("/api/upload-dsr", tags=["ingestion"])
@router.post("/api/upload-report", tags=["ingestion"])
async def upload_dsr_workbook(
    file: UploadFile = File(..., description="DSR Report file (.xlsx, .xlsm, .csv, .txt, .tsv)"),
    period: str | None = Form(None, description="Optional period label e.g. AUG2026, SEP2026"),
    uploaded_by: str = Form("Reporting Agent", description="Name of agent or manager uploading"),
    table_type: str | None = Form(None, description="Optional target table: auto, booking, lead, vehicle"),
    mode: str = Form("replace", description="replace the month, or append to it"),
):
    """
    Ingest a DSR report file (Excel workbook, CSV, or Text format).
    Rebuilds/updates facts (leads, bookings, vehicles, test drives),
    updates period alignment in Supabase, and notifies connected web clients.

    The dashboard sends Excel workbooks to /api/upload-report/start instead,
    because a workbook takes longer than a gateway will hold a request open.
    """
    mode = mode if mode in ("replace", "append") else "replace"
    fname = file.filename or "unknown_report.txt"
    fname_lower = fname.lower()
    allowed_exts = (".xlsx", ".xlsm", ".xls", ".csv", ".txt", ".tsv")
    if not fname_lower.endswith(allowed_exts):
        raise HTTPException(
            400,
            f"Invalid file format. Please upload an Excel (.xlsx, .xlsm), CSV (.csv), or Text (.txt, .tsv) file."
        )

    content = await file.read()
    if not content:
        raise HTTPException(400, "Uploaded file is empty.")

    period_label, period_start, period_end = _resolve_period(period, fname)

    # 1. Handle Plain Text, CSV, and TSV files
    if fname_lower.endswith((".csv", ".txt", ".tsv")):
        from etl.text_parser import ingest_text_report
        try:
            parse_result = ingest_text_report(
                content=content,
                filename=fname,
                period_label=period_label,
                period_start=period_start,
                period_end=period_end,
                uploaded_by=uploaded_by,
                table_type=table_type,
                mode=mode,
            )
            counts = parse_result.get("counts", {})
            warnings = parse_result.get("warnings", [])

            if is_db_ready():
                try:
                    with db_connect(autocommit=True) as cx:
                        cx.execute("""
                            INSERT INTO etl_run (source_file, file_modified, finished_at, row_counts, notes)
                            VALUES (%s, now(), now(), %s, %s)
                        """, (
                            fname,
                            json.dumps(counts),
                            f"Uploaded {parse_result.get('detected_format', 'text')} by {uploaded_by}",
                        ))
                except Exception:
                    pass

            try:
                await broker.notify("workbook_reload", "etl_run", {"counts": counts, "period": period_label})
            except Exception:
                pass

            return {
                "status": "success",
                "message": f"Successfully ingested {parse_result.get('detected_format', 'report')} '{fname}' for {period_label}.",
                "filename": fname,
                "file_type": "text/csv",
                "detected_format": parse_result.get("detected_format"),
                "detected_table": parse_result.get("detected_table"),
                # A single-table CSV is always added to the month, whatever
                # was asked for - text_parser reports what it actually did.
                "mode": parse_result.get("mode", mode),
                "period": period_label,
                "period_range": f"{period_start} to {period_end}",
                "uploaded_by": uploaded_by,
                "counts": counts,
                "warnings": warnings,
                "environment": "postgres" if is_db_ready() else "fallback",
            }
        except Exception as exc:
            import traceback
            traceback.print_exc()
            raise HTTPException(400, f"Failed to ingest CSV/TXT report: {exc}")

    # 2. Handle Excel Workbooks (.xlsx, .xlsm)
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as exc:
        raise HTTPException(400, f"Could not parse Excel workbook: {exc}")
    problem = workbook_problem(wb)
    if problem:
        raise HTTPException(400, problem)

    if is_db_ready():
        try:
            with db_connect(autocommit=True) as cx:
                # FIX (2026-10-08): one transaction around the whole load, the
                # same as _ingest_worker. On a bare autocommit connection a
                # replace that failed part-way had already committed reset()'s
                # deletes, leaving the month empty.
                with cx.transaction():
                    loader = Loader(cx, wb, period_label, period_start, period_end, mode=mode)
                    counts = loader.run()
                    cx.execute("""
                        INSERT INTO etl_run (source_file, file_modified, finished_at, row_counts, notes)
                        VALUES (%s, now(), now(), %s, %s)
                    """, (
                        fname,
                        json.dumps(counts),
                        f"Uploaded by {uploaded_by} ({mode})"
                        + (("; " + "; ".join(loader.warnings)) if loader.warnings else ""),
                    ))

            try:
                await broker.notify("workbook_reload", "etl_run", {"counts": counts, "period": period_label})
            except Exception:
                pass

            return {
                "status": "success",
                "message": f"Successfully ingested '{fname}' into PostgreSQL database for {period_label}.",
                "filename": fname,
                "file_type": "excel",
                "mode": mode,
                "period": period_label,
                "period_range": f"{period_start} to {period_end}",
                "uploaded_by": uploaded_by,
                "counts": counts,
                "warnings": loader.warnings,
                "environment": "postgres",
            }
        except Exception as exc:
            import traceback
            tb = traceback.format_exc()
            print("ERROR IN LOADER:", tb)
            raise HTTPException(500, f"Database ingestion failed: {exc}")

    # Fallback mode (when running in serverless / offline without DB):
    counts = {}
    sheet_map = {
        "Live Booking": "booking",
        "Vehicle Status": "vehicle",
        "Enquiries": "lead",
        "Enquiry": "lead",
        "Test Drive": "test_drive",
        "Allotment": "allotment",
        "Reg Report": "registration",
    }
    for sname in wb.sheetnames:
        for prefix, table in sheet_map.items():
            if prefix.lower() in sname.lower():
                ws = wb[sname]
                filled_rows = sum(
                    1 for r in range(2, min(ws.max_row + 1, 5000))
                    if ws.cell(r, 1).value or ws.cell(r, 2).value
                )
                counts[table] = max(counts.get(table, 0), filled_rows)

    return {
        "status": "success",
        "message": f"Successfully parsed '{fname}' for {period_label} (Serverless/Preview mode).",
        "filename": fname,
        "period": period_label,
        "period_range": f"{period_start} to {period_end}",
        "uploaded_by": uploaded_by,
        "counts": counts,
        "warnings": [],
        "environment": "serverless",
    }


@router.get("/api/etl-history", tags=["ingestion"])
def etl_history(limit: int = Query(10, le=50)):
    """Return recent Excel workbook ingestion runs for audit and status."""
    if is_db_ready():
        try:
            return fetch_all("""
                SELECT run_id, source_file, file_modified, finished_at, row_counts, notes
                FROM etl_run
                ORDER BY finished_at DESC
                LIMIT %s
            """, (limit,))
        except Exception:
            pass
    return [
        {
            "run_id": 1,
            "source_file": "DSR August 2026.xlsx",
            "file_modified": datetime.now().isoformat(),
            "finished_at": datetime.now().isoformat(),
            "row_counts": {"vehicle": 107, "lead": 2154, "booking": 133, "allotment": 20, "registration": 21},
            "notes": "Initial seed load",
        }
    ]
